"""Отчёт по валидации: make eval

Считает всё, что ТЗ просит в отчёте: метрики классификации, отсутствие
галлюцинаций, близость черновиков к реальным ответам админов, долю
«нужен человек», стоимость на тикет и задержку (p50/p95).

Черновики через LLM считаются на случайной выборке (--llm-sample), чтобы не
упираться в лимиты бесплатного ключа; остальные — в офлайн-режиме. Стоимость
на тикет считается с учётом доли обращений, которым LLM не нужен (шаблоны).

Точность классификации меряется двумя способами:
  - против слабых меток из исходов админов (есть у всех тикетов валидации);
  - против ручной разметки команды (data/gold/label_me.xlsx), если она заполнена.
В отчёт попадают только агрегированные цифры: данные Grand Mobile не публикуются.
"""

import argparse
import json
import random
import statistics
import sys
import time
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from sklearn.feature_extraction.text import TfidfVectorizer  # noqa: E402
from sklearn.metrics import accuracy_score, f1_score  # noqa: E402

from copilot import config, llm  # noqa: E402
from copilot.labels import weak_label  # noqa: E402
from copilot.pipeline import triage_full  # noqa: E402
from copilot.predict import ticket_to_request  # noqa: E402
from copilot.rules import is_valid_ref, parse_refs  # noqa: E402

GOLD_PATH = ROOT / "data" / "gold" / "label_me.xlsx"


def load_gold() -> dict[str, dict]:
    if not GOLD_PATH.exists():
        return {}
    from openpyxl import load_workbook
    ws = load_workbook(GOLD_PATH, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h) for h in rows[0]]
    gold = {}
    for r in rows[1:]:
        rec = dict(zip(head, r))
        if rec.get("category") and rec.get("priority"):
            gold[rec["ticket_id"]] = {"category": str(rec["category"]).strip(),
                                      "priority": str(rec["priority"]).strip(),
                                      "needs_human": str(rec.get("needs_human") or "").strip().lower() == "да"}
    return gold


def pct(x: float) -> str:
    return f"{x:.1%}"


def confusion(pairs: list[tuple[str, str]]) -> dict[str, dict[str, int]]:
    """{истинная категория: {предсказанная: число}} — только ненулевые клетки."""
    table: dict[str, Counter] = {}
    for gold, pred in pairs:
        table.setdefault(gold, Counter())[pred] += 1
    return {g: dict(c.most_common()) for g, c in sorted(table.items())}


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--llm-sample", type=int, default=100, help="сколько тикетов прогнать через LLM")
    ap.add_argument("--llm-rpm", type=float, default=12, help="лимит запросов к LLM в минуту (бесплатный ключ)")
    ap.add_argument("--out", default=str(config.ARTIFACTS_DIR / "report.json"))
    args = ap.parse_args()

    tickets = [json.loads(line) for line in open(config.VALID_PATH, encoding="utf-8")]
    random.seed(42)
    llm_ids = set()
    if llm.enabled() and args.llm_sample:
        llm_ids = {t["ticket_id"] for t in random.sample(tickets, min(args.llm_sample, len(tickets)))}
    mode = config.LLM_MODE

    rows = []
    last_llm = 0.0
    for t in tickets:
        config.LLM_MODE = mode if t["ticket_id"] in llm_ids else "offline"
        if t["ticket_id"] in llm_ids and args.llm_rpm:
            # Пауза вне замера задержки: бесплатный ключ ограничен по запросам в минуту.
            time.sleep(max(0.0, last_llm + 60 / args.llm_rpm - time.perf_counter()))
            last_llm = time.perf_counter()
        started = time.perf_counter()
        pred, meta = triage_full(ticket_to_request(t))
        rows.append({"t": t, "pred": pred, "meta": meta, "latency": time.perf_counter() - started,
                     "llm_sampled": t["ticket_id"] in llm_ids})
    config.LLM_MODE = mode

    report: dict = {"n_tickets": len(rows), "llm_model": config.LLM_MODEL if llm_ids else None}

    # --- Классификация против слабых меток (исходы админов) ---
    for kind in ("complaint", "report"):
        pairs = [(weak_label(r["t"]), r["pred"]["category"]) for r in rows if r["t"]["kind"] == kind]
        pairs = [(g, p) for g, p in pairs if g]
        report[f"category_vs_outcome_{kind}"] = {
            "n": len(pairs), "accuracy": round(accuracy_score(*zip(*pairs)), 3) if pairs else None,
            "confusion": confusion(pairs)}

    # --- Классификация против ручной разметки команды ---
    gold = load_gold()
    by_id = {r["t"]["ticket_id"]: r["pred"] for r in rows}
    g = [(gold[i], by_id[i]) for i in gold if i in by_id]
    if g:
        report["gold"] = {
            "n": len(g),
            "category_accuracy": round(accuracy_score([x["category"] for x, _ in g], [p["category"] for _, p in g]), 3),
            "priority_macro_f1": round(f1_score([x["priority"] for x, _ in g], [p["priority"] for _, p in g],
                                                average="macro", zero_division=0), 3),
            "needs_human_accuracy": round(accuracy_score([x["needs_human"] for x, _ in g],
                                                         [p["needs_human"] for _, p in g]), 3),
        }

    # --- Галлюцинации: каждая цитата и каждый упомянутый пункт должны существовать ---
    cited = [c for r in rows for c in r["pred"]["draft_reply"]["cited"]]
    mentioned = [m for r in rows for m in parse_refs(r["pred"]["draft_reply"]["text"])]
    report["hallucinations"] = {
        "citations_total": len(cited),
        "citations_invalid": sum(not is_valid_ref(c) for c in cited),
        "mentioned_refs_invalid": sum(not is_valid_ref(m) for m in mentioned),
        "drafts_with_citation": pct(sum(bool(r["pred"]["draft_reply"]["cited"]) for r in rows) / len(rows)),
    }

    # --- Близость черновика к реальному ответу админа (прокси для слепой оценки) ---
    def admin_reply(t):
        return " ".join(m["text"] for m in t["messages"] if m["role"] == "admin")

    pairs = [(r["pred"]["draft_reply"]["text"], admin_reply(r["t"]), r) for r in rows if admin_reply(r["t"])]
    vec = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5)).fit([a for a, _, _ in pairs] + [b for _, b, _ in pairs])
    sims = []
    for draft, real, r in pairs:
        a, b = vec.transform([draft, real])
        sims.append((float((a @ b.T).toarray()[0, 0]), r))
    by_source: dict[str, list] = {}
    for s, r in sims:
        by_source.setdefault(r["meta"].get("draft_source") or "none", []).append(s)
    report["draft_similarity_to_admin"] = {
        "mean_all": round(statistics.mean(s for s, _ in sims), 3),
        "by_source": {k: {"n": len(v), "mean": round(statistics.mean(v), 3)} for k, v in by_source.items()},
    }

    # --- Язык ответа для kk/uz ---
    lang_rows = [r for r in rows if r["pred"]["lang"] in ("kk", "uz") and r["llm_sampled"]]
    report["kk_uz"] = {
        "tickets_kk_uz": sum(r["pred"]["lang"] in ("kk", "uz") for r in rows),
        "llm_reply_same_language": (pct(sum(r["pred"]["draft_reply"]["lang"] == r["pred"]["lang"] for r in lang_rows)
                                        / len(lang_rows)) if lang_rows else None),
    }

    # --- Доли, стоимость, задержка ---
    report["shares"] = {
        "needs_human": pct(sum(r["pred"]["needs_human"] for r in rows) / len(rows)),
        "spam_or_empty": pct(sum(r["pred"]["is_spam_or_duplicate"] for r in rows) / len(rows)),
        "category": dict(Counter(r["pred"]["category"] for r in rows).most_common()),
        "priority": dict(sorted(Counter(r["pred"]["priority"] for r in rows).items())),
    }
    sampled = [r for r in rows if r["llm_sampled"]]
    llm_rows = [r for r in sampled if r["meta"].get("draft_source") == "llm"]
    costs = [r["meta"]["cost_usd"] for r in llm_rows]
    if sampled:
        llm_share = len(llm_rows) / len(sampled)
        report["cost"] = {
            "sampled_tickets": len(sampled),
            "llm_used_share": pct(llm_share),
            "llm_errors_fallback_offline": sum("llm_error" in r["meta"] for r in sampled),
            "llm_error_types": dict(Counter(r["meta"]["llm_error"][:60] for r in sampled if "llm_error" in r["meta"])),
            "usd_per_llm_call_mean": round(statistics.mean(costs), 6) if costs else None,
            "usd_per_llm_call_max": round(max(costs), 6) if costs else None,
            "usd_per_ticket_blended": round(statistics.mean(costs) * llm_share, 6) if costs else 0,
            "tokens_in_mean": round(statistics.mean(r["meta"]["tokens_in"] for r in llm_rows)) if llm_rows else None,
            "tokens_out_mean": round(statistics.mean(r["meta"]["tokens_out"] for r in llm_rows)) if llm_rows else None,
            "limit_usd": 0.002,
        }

    def lat(xs):
        xs = sorted(xs)
        return {"p50_s": round(xs[len(xs) // 2], 3), "p95_s": round(xs[int(len(xs) * 0.95) - 1], 3)} if xs else None

    report["latency"] = {
        "offline": lat([r["latency"] for r in rows if not r["llm_sampled"]]),
        "with_llm_sample": lat([r["latency"] for r in sampled]),
        "limit_p95_s": 5.0,
    }

    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    Path(args.out).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
